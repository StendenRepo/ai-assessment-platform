from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

_MAX_CRITERIA = 300
_FALLBACK_MIN_WORDS = 3
_FALLBACK_MAX_CHARS = 400
_HEADER_HINTS = (
    "criteri",
    "indicator",
    "description",
    "outcome",
    "competen",
    "learning goal",
    "rubric",
)


@dataclass(frozen=True)
class Criterion:
    key: str
    text: str


def extract_criteria(
    rubric_path: Path | None,
    ext: str | None,
    fallback_text: str | None,
) -> list[Criterion]:
    criteria: list[Criterion] = []
    if ext == ".xlsx" and rubric_path is not None and rubric_path.exists():
        criteria = _from_xlsx(rubric_path)
    if not criteria:
        criteria = _from_text(fallback_text)
    return _dedupe_keys(criteria[:_MAX_CRITERIA])


def _cell_str(value) -> str:
    return "" if value is None else str(value).strip()


def _looks_like_header(cells: list[str]) -> bool:
    joined = " ".join(cells).lower()
    return any(hint in joined for hint in _HEADER_HINTS)


def _from_xlsx(path: Path) -> list[Criterion]:
    try:
        workbook = load_workbook(filename=path, read_only=True, data_only=True)
    except Exception:
        return []
    try:
        sheet = workbook.active
        if sheet is None:
            return []

        criteria: list[Criterion] = []
        for index, row in enumerate(sheet.iter_rows(values_only=True)):
            cells = [_cell_str(value) for value in row]
            non_empty = [c for c in cells if c]
            if not non_empty:
                continue
            if index == 0 and _looks_like_header(non_empty):
                continue
            criteria.append(Criterion(key=non_empty[0], text=" ".join(non_empty)))
        return criteria
    finally:
        workbook.close()


def _from_text(text: str | None) -> list[Criterion]:
    if not text:
        return []
    criteria: list[Criterion] = []
    for raw_line in text.splitlines():
        line = raw_line.strip(" \t-*•.")
        if len(line) > _FALLBACK_MAX_CHARS:
            continue
        if len(line.split()) < _FALLBACK_MIN_WORDS:
            continue
        criteria.append(Criterion(key=line[:120], text=line))
    return criteria


def _dedupe_keys(criteria: list[Criterion]) -> list[Criterion]:
    seen: dict[str, int] = {}
    result: list[Criterion] = []
    for criterion in criteria:
        count = seen.get(criterion.key, 0) + 1
        seen[criterion.key] = count
        key = criterion.key if count == 1 else f"{criterion.key} ({count})"
        result.append(Criterion(key=key, text=criterion.text))
    return result
