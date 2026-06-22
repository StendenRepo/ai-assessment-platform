"""Parsing of uploaded student lists (.xlsx / .csv).

This module is intentionally free of any database or FastAPI concerns so it can
be unit-tested in isolation. It turns an uploaded file into a list of
``ParsedRow`` objects (one per non-empty data row) and raises
``ImportParseError`` for problems with the file as a whole (wrong type, no
recognisable header row, or missing required columns).
"""
import csv
import io
from dataclasses import dataclass
from typing import List

# Accepted header spellings, normalised (lowercased, stripped of spaces/_/-).
_NAME_HEADERS = {"name", "studentname", "fullname", "naam", "student"}
_NUMBER_HEADERS = {
    "studentnumber",
    "studentno",
    "studentnr",
    "studentid",
    "number",
    "nummer",
    "no",
    "id",
}


class ImportParseError(Exception):
    """Raised when the file cannot be parsed at all (vs. a single bad row)."""


@dataclass
class ParsedRow:
    row_number: int  # 1-based row number in the original sheet (header is row 1)
    name: str
    student_number: str


def _normalise(header: object) -> str:
    return "".join(str(header or "").lower().split()).replace("_", "").replace("-", "")


def _grid_from_xlsx(data: bytes) -> List[List[str]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        # openpyxl raises its own errors (InvalidFileException, BadZipFile, ...)
        # on a corrupt or non-xlsx file; surface a clean user-facing message.
        raise ImportParseError(
            "The Excel file could not be read. Make sure it is a valid, "
            "uncorrupted .xlsx file."
        )
    try:
        sheet = workbook.active
        return [
            ["" if cell is None else str(cell) for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()


def _grid_from_csv(data: bytes) -> List[List[str]]:
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    # EU spreadsheets often export CSV with ';'. Pick whichever appears more.
    delimiter = ";" if sample.count(";") > sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    return [["" if cell is None else str(cell) for cell in row] for row in reader]


def parse_student_file(filename: str, data: bytes) -> List[ParsedRow]:
    """Parse an uploaded students file into rows.

    Raises ImportParseError for an unsupported file type, an empty file, or a
    header row that is missing a name and/or student-number column.
    """
    lowered = (filename or "").lower()
    if lowered.endswith(".xlsx"):
        grid = _grid_from_xlsx(data)
    elif lowered.endswith(".csv"):
        grid = _grid_from_csv(data)
    else:
        raise ImportParseError(
            "Unsupported file type. Please upload a .xlsx or .csv file."
        )

    # Find the first non-empty row; treat it as the header.
    header_index = next(
        (i for i, row in enumerate(grid) if any(str(c).strip() for c in row)),
        None,
    )
    if header_index is None:
        raise ImportParseError("The file is empty.")

    header = grid[header_index]
    name_col = number_col = None
    for col, cell in enumerate(header):
        key = _normalise(cell)
        if name_col is None and key in _NAME_HEADERS:
            name_col = col
        elif number_col is None and key in _NUMBER_HEADERS:
            number_col = col

    if name_col is None or number_col is None:
        raise ImportParseError(
            "Could not find the required columns. The sheet needs a 'Name' "
            "column and a 'Student Number' column in the first row."
        )

    rows: List[ParsedRow] = []
    for offset, row in enumerate(grid[header_index + 1 :]):
        if not any(str(c).strip() for c in row):
            continue  # skip fully blank rows (common trailing rows)
        name = row[name_col].strip() if name_col < len(row) else ""
        number = row[number_col].strip() if number_col < len(row) else ""
        rows.append(
            ParsedRow(
                row_number=header_index + 2 + offset,
                name=name,
                student_number=number,
            )
        )
    return rows
