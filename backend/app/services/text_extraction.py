"""Plain-text extraction for uploaded documents (evidence, rubrics, module books).

``extract_document_text`` is best-effort (returns ``None`` on failure) for module
documents where a failed parse must not block upload.

``read_stored_evidence_text`` reads an evidence file from disk and returns text
for overlap detection, AI grounding, and content APIs — including PDF, Word,
Excel, and markdown uploads.
"""

from __future__ import annotations

import csv
import io
from pathlib import Path

from app.models.evidence import Evidence
from app.models.enums import FileType

# Extensions accepted for student evidence upload (key includes leading dot).
EVIDENCE_EXTENSIONS: dict[str, FileType] = {
    ".md": FileType.markdown,
    ".txt": FileType.other,
    ".pdf": FileType.pdf,
    ".docx": FileType.docx,
    ".xlsx": FileType.xlsx,
    ".csv": FileType.other,
}


def evidence_extension(filename: str) -> str:
    return Path(filename or "").suffix.lower()


def is_supported_evidence_filename(filename: str) -> bool:
    return evidence_extension(filename) in EVIDENCE_EXTENSIONS


def resolve_evidence_file_type(filename: str) -> FileType:
    ext = evidence_extension(filename)
    file_type = EVIDENCE_EXTENSIONS.get(ext)
    if file_type is None:
        allowed = ", ".join(sorted(EVIDENCE_EXTENSIONS))
        raise ValueError(f"Unsupported file type '{ext}'. Allowed: {allowed}")
    return file_type


def extract_document_text(raw: bytes, ext: str) -> str | None:
    """Return plain text extracted from *raw* for file extension *ext*.

    *ext* is the lowercased extension including the leading dot (e.g. ".pdf").
    Returns ``None`` if the type is unsupported, the file cannot be parsed, or
    the result is empty after stripping. Never raises.
    """
    try:
        if ext == ".pdf":
            text = _from_pdf(raw)
        elif ext == ".docx":
            text = _from_docx(raw)
        elif ext == ".xlsx":
            text = _from_xlsx(raw)
        elif ext == ".csv":
            text = _from_csv(raw)
        elif ext in (".md", ".txt"):
            text = raw.decode("utf-8", errors="ignore")
        else:
            return None
    except Exception:
        return None

    text = (text or "").strip()
    return text or None


def extract_document_text_strict(raw: bytes, filename: str) -> str:
    """Like ``extract_document_text`` but raises ``ValueError`` on failure."""
    ext = evidence_extension(filename)
    resolve_evidence_file_type(filename)  # raises if unsupported
    text = extract_document_text(raw, ext)
    if not text:
        raise ValueError(f"Could not extract text from '{filename}'")
    return text


def read_stored_evidence_text(evidence: Evidence, upload_dir: Path) -> str:
    """Read extractable plain text from a stored evidence file."""
    path = upload_dir / evidence.file_path
    if not path.exists():
        return ""
    try:
        raw = path.read_bytes()
    except OSError:
        return ""
    if not raw:
        return ""

    ext = evidence_extension(evidence.file_name) or path.suffix.lower()

    # Upload pipeline stores extracted UTF-8 text on disk (any source format).
    try:
        decoded = raw.decode("utf-8")
        if decoded.strip():
            return decoded
    except UnicodeDecodeError:
        pass

    text = extract_document_text(raw, ext)
    if text:
        return text

    stored_ext = path.suffix.lower()
    if stored_ext and stored_ext != ext:
        text = extract_document_text(raw, stored_ext)
        if text:
            return text
    return ""


def _from_pdf(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    return "\n\n".join(page.extract_text() or "" for page in reader.pages)


def _from_docx(raw: bytes) -> str:
    from docx import Document as DocxDocument

    doc = DocxDocument(io.BytesIO(raw))
    return "\n".join(paragraph.text for paragraph in doc.paragraphs)


def _from_xlsx(raw: bytes) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        for row in sheet.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                lines.append(" ".join(cells))
    wb.close()
    return "\n".join(lines)


def _from_csv(raw: bytes) -> str:
    decoded = raw.decode("utf-8-sig", errors="ignore")
    reader = csv.reader(io.StringIO(decoded))
    lines: list[str] = []
    for row in reader:
        cells = [cell.strip() for cell in row if cell and cell.strip()]
        if cells:
            lines.append(" ".join(cells))
    return "\n".join(lines)
