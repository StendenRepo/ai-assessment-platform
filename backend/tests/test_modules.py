import uuid
import pytest

from app.core.security import hash_password

LOGIN_URL = "/api/v1/auth/login"
MODULES_URL = "/api/v1/modules"

_TEACHER_EMAIL = "modules_teacher@test.com"
_TEACHER_PASSWORD = "password123"


# ---------------------------------------------------------------------------
# Module-local teacher fixture (avoids email collision with other test files)
# ---------------------------------------------------------------------------

@pytest.fixture
def modules_teacher(db):
    """Get-or-create a teacher used exclusively by this test module.

    The teacher is intentionally NOT deleted in teardown: the session-scoped
    SQLite DB is dropped at the end of the test session anyway, and deleting
    the teacher would cascade-null the teacher_id FK on any modules created
    during the test (violating the NOT NULL constraint).
    """
    from app.models.teacher import Teacher
    from app.core.security import hash_password

    existing = db.query(Teacher).filter(Teacher.email == _TEACHER_EMAIL).first()
    if existing:
        yield existing
        return

    t = Teacher(
        id=uuid.uuid4(),
        name="Modules Teacher",
        email=_TEACHER_EMAIL,
        password_hash=hash_password(_TEACHER_PASSWORD),
    )
    db.add(t)
    db.commit()
    db.refresh(t)
    yield t
    # No teardown delete — session-scoped DB is dropped after all tests finish.


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _login(client, email: str = _TEACHER_EMAIL, password: str = _TEACHER_PASSWORD) -> str:
    """Return a raw access token. Defaults to the modules_teacher credentials."""
    res = client.post(LOGIN_URL, json={"email": email, "password": password})
    assert res.status_code == 200, res.text
    return res.json()["access_token"]


def _auth(client, email: str = _TEACHER_EMAIL, password: str = _TEACHER_PASSWORD) -> dict:
    """Return an Authorization header dict."""
    return {"Authorization": f"Bearer {_login(client, email, password)}"}


def _create_module(client, headers, name="Test Module", academic_year="2024-2025"):
    res = client.post(
        MODULES_URL,
        json={"name": name, "academic_year": academic_year},
        headers=headers,
    )
    assert res.status_code == 201, res.text
    return res.json()


# ---------------------------------------------------------------------------
# PATCH /modules/{module_id} — rename
# ---------------------------------------------------------------------------

class TestRenameModule:
    def test_rename_updates_name(self, client, modules_teacher):
        headers = _auth(client)
        module = _create_module(client, headers)

        res = client.patch(
            f"{MODULES_URL}/{module['id']}",
            json={"name": "Renamed Module"},
            headers=headers,
        )
        assert res.status_code == 200
        body = res.json()
        assert body["name"] == "Renamed Module"
        assert body["id"] == module["id"]

    def test_rename_preserves_academic_year(self, client, modules_teacher):
        headers = _auth(client)
        module = _create_module(client, headers, academic_year="2023-2024")

        res = client.patch(
            f"{MODULES_URL}/{module['id']}",
            json={"name": "New Name"},
            headers=headers,
        )
        assert res.status_code == 200
        assert res.json()["academic_year"] == "2023-2024"

    def test_rename_can_also_update_academic_year(self, client, modules_teacher):
        headers = _auth(client)
        module = _create_module(client, headers, academic_year="2023-2024")

        res = client.patch(
            f"{MODULES_URL}/{module['id']}",
            json={"name": "Updated", "academic_year": "2025-2026"},
            headers=headers,
        )
        assert res.status_code == 200
        body = res.json()
        assert body["name"] == "Updated"
        assert body["academic_year"] == "2025-2026"

    def test_rename_empty_name_returns_422(self, client, modules_teacher):
        headers = _auth(client)
        module = _create_module(client, headers)

        res = client.patch(
            f"{MODULES_URL}/{module['id']}",
            json={"name": "   "},
            headers=headers,
        )
        assert res.status_code == 422

    def test_rename_missing_name_returns_422(self, client, modules_teacher):
        headers = _auth(client)
        module = _create_module(client, headers)

        res = client.patch(
            f"{MODULES_URL}/{module['id']}",
            json={},
            headers=headers,
        )
        assert res.status_code == 422

    def test_rename_nonexistent_module_returns_404(self, client, modules_teacher):
        headers = _auth(client)
        fake_id = str(uuid.uuid4())

        res = client.patch(
            f"{MODULES_URL}/{fake_id}",
            json={"name": "Ghost"},
            headers=headers,
        )
        assert res.status_code == 404

    def test_rename_other_teachers_module_returns_404(self, client, modules_teacher, db):
        """A teacher cannot rename a module they don't own."""
        from app.models.module import Module
        from app.core.security import hash_password
        from app.models.teacher import Teacher

        other_email = f"other_rename_{uuid.uuid4().hex[:8]}@test.com"
        other = Teacher(
            id=uuid.uuid4(),
            name="Other Teacher",
            email=other_email,
            password_hash=hash_password("password123"),
        )
        db.add(other)
        db.commit()

        other_module = Module(
            id=uuid.uuid4(),
            teacher_id=other.id,
            name="Other's Module",
        )
        db.add(other_module)
        db.commit()

        # Log in as the modules teacher and try to rename the other's module
        headers = _auth(client)
        res = client.patch(
            f"{MODULES_URL}/{other_module.id}",
            json={"name": "Stolen Name"},
            headers=headers,
        )
        assert res.status_code == 404

        db.delete(other_module)
        db.delete(other)
        db.commit()

    def test_rename_without_auth_returns_401(self, client, modules_teacher):
        headers = _auth(client)
        module = _create_module(client, headers)

        res = client.patch(
            f"{MODULES_URL}/{module['id']}",
            json={"name": "No Auth"},
        )
        assert res.status_code == 401


# ---------------------------------------------------------------------------
# DELETE /modules/{module_id}
# ---------------------------------------------------------------------------

class TestDeleteModule:
    def test_delete_returns_204(self, client, modules_teacher):
        headers = _auth(client)
        module = _create_module(client, headers, name="To Delete")

        res = client.delete(f"{MODULES_URL}/{module['id']}", headers=headers)
        assert res.status_code == 204

    def test_deleted_module_no_longer_in_list(self, client, modules_teacher):
        headers = _auth(client)
        module = _create_module(client, headers, name="Will Be Gone")

        client.delete(f"{MODULES_URL}/{module['id']}", headers=headers)

        list_res = client.get(MODULES_URL, headers=headers)
        assert list_res.status_code == 200
        ids = [m["id"] for m in list_res.json()]
        assert module["id"] not in ids

    def test_delete_nonexistent_module_returns_404(self, client, modules_teacher):
        headers = _auth(client)
        fake_id = str(uuid.uuid4())

        res = client.delete(f"{MODULES_URL}/{fake_id}", headers=headers)
        assert res.status_code == 404

    def test_delete_other_teachers_module_returns_404(self, client, modules_teacher, db):
        """A teacher cannot delete a module they don't own."""
        from app.models.module import Module
        from app.core.security import hash_password
        from app.models.teacher import Teacher

        other_email = f"other_delete_{uuid.uuid4().hex[:8]}@test.com"
        other = Teacher(
            id=uuid.uuid4(),
            name="Other Teacher 2",
            email=other_email,
            password_hash=hash_password("password123"),
        )
        db.add(other)
        db.commit()

        other_module = Module(
            id=uuid.uuid4(),
            teacher_id=other.id,
            name="Other's Module 2",
        )
        db.add(other_module)
        db.commit()

        headers = _auth(client)
        res = client.delete(f"{MODULES_URL}/{other_module.id}", headers=headers)
        assert res.status_code == 404

        db.delete(other_module)
        db.delete(other)
        db.commit()

    def test_delete_without_auth_returns_401(self, client, modules_teacher):
        headers = _auth(client)
        module = _create_module(client, headers, name="Auth Required")

        res = client.delete(f"{MODULES_URL}/{module['id']}")
        assert res.status_code == 401

    def test_get_after_delete_returns_404(self, client, modules_teacher):
        headers = _auth(client)
        module = _create_module(client, headers, name="Check After Delete")

        client.delete(f"{MODULES_URL}/{module['id']}", headers=headers)

        get_res = client.get(f"{MODULES_URL}/{module['id']}", headers=headers)
        assert get_res.status_code == 404

    def test_delete_removes_evidence_files_from_disk(self, client, modules_teacher, db, tmp_path, monkeypatch):
        """Evidence files on disk must be removed when the module is deleted."""
        from app.models.project import Project
        from app.models.student import Student
        from app.models.evidence import Evidence
        from app.models.enums import EmbeddingStatus, FileType, SourceType
        import app.config as app_config

        # ── Redirect UPLOAD_DIR to tmp_path so we don't touch real storage ──
        # The delete_module endpoint computes:
        #   EVIDENCE_UPLOAD_DIR = Path(settings.UPLOAD_DIR) / "evidence"
        # Patching settings.UPLOAD_DIR makes it resolve to our tmp dir.
        monkeypatch.setattr(app_config.settings, "UPLOAD_DIR", str(tmp_path))

        fake_evidence_dir = tmp_path / "evidence"
        fake_evidence_dir.mkdir(parents=True, exist_ok=True)

        headers = _auth(client)
        module = _create_module(client, headers, name="Module With Files")

        # Create a group + student + evidence record with a real file on disk
        project = Project(
            id=uuid.uuid4(),
            module_id=uuid.UUID(module["id"]),
            name="Group A",
            group_name="Group A",
        )
        db.add(project)
        db.commit()

        student = Student(
            name="Test Student",
            student_number="999",
        )
        db.add(student)
        db.flush()
        student.projects.append(project)
        db.commit()

        # Write a fake evidence file to disk
        student_dir = fake_evidence_dir / str(student.student_number)
        student_dir.mkdir(parents=True)
        fake_file = student_dir / "abc123_report.md"
        fake_file.write_text("# Evidence content")

        relative_path = str(fake_file.relative_to(fake_evidence_dir))
        ev = Evidence(
            id=uuid.uuid4(),
            student_id=student.student_number,
            file_name="report.md",
            file_type=FileType.markdown,
            file_path=relative_path,
            source_type=SourceType.upload,
            embedding_status=EmbeddingStatus.completed,
        )
        db.add(ev)
        db.commit()

        assert fake_file.exists(), "Precondition: file must exist before delete"

        # Delete the module via the API
        res = client.delete(f"{MODULES_URL}/{module['id']}", headers=headers)
        assert res.status_code == 204

        assert not fake_file.exists(), "Evidence file must be removed from disk after module delete"

def test_admin_cannot_upload_rubric_for_other_teachers_module(client, db, teacher):
    """Admins are blocked (403) from uploading a rubric to a module they don't own."""
    from app.models.module import Module
    from app.models.teacher import Teacher

    admin = Teacher(
        id=uuid.uuid4(),
        name="Admin Teacher",
        email="admin@test.com",
        password_hash=hash_password("password123"),
        is_admin=True,
    )
    db.add(admin)
    db.commit()
    db.refresh(admin)

    module = Module(id=uuid.uuid4(), teacher_id=teacher.id, name="Owned by non-admin")
    db.add(module)
    db.commit()
    db.refresh(module)

    try:
        headers = _auth(client, "admin@test.com", "password123")

        upload = client.post(
            f"{MODULES_URL}/{module.id}/rubric",
            files={"file": ("rubric.pdf", b"%PDF-1.4 test rubric", "application/pdf")},
            headers=headers,
        )
        assert upload.status_code == 403
    finally:
        db.query(Module).filter(Module.id == module.id).delete()
        db.query(Teacher).filter(Teacher.id == admin.id).delete()
        db.commit()


def test_teacher_can_upload_replace_and_delete_module_book(client, db, teacher):
    from app.models.file_record import FileRecord
    from app.models.module import Module

    module = Module(id=uuid.uuid4(), teacher_id=teacher.id, name="With module book")
    db.add(module)
    db.commit()
    db.refresh(module)

    try:
        headers = _auth(client, "teacher@test.com", "password123")

        upload = client.post(
            f"{MODULES_URL}/{module.id}/module-book",
            files={"file": ("book.pdf", b"%PDF-1.4 module book", "application/pdf")},
            headers=headers,
        )
        assert upload.status_code == 200
        body = upload.json()
        assert body["id"] == str(module.id)
        assert body["module_book_file"] is not None
        assert body["module_book_file"]["file_name"] == "book.pdf"

        # Replacing swaps the linked file and removes the previous record.
        replace = client.post(
            f"{MODULES_URL}/{module.id}/module-book",
            files={
                "file": (
                    "book.docx",
                    b"PK\x03\x04 module book v2",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            },
            headers=headers,
        )
        assert replace.status_code == 200
        assert replace.json()["module_book_file"]["file_name"] == "book.docx"
        assert db.query(FileRecord).count() == 1

        remove = client.delete(f"{MODULES_URL}/{module.id}/module-book", headers=headers)
        assert remove.status_code == 204

        db.refresh(module)
        assert module.module_book_id is None
        assert db.query(FileRecord).count() == 0
    finally:
        db.query(FileRecord).delete()
        db.query(Module).filter(Module.id == module.id).delete()
        db.commit()


def test_module_book_rejects_unsupported_extension(client, db, teacher):
    from app.models.module import Module

    module = Module(id=uuid.uuid4(), teacher_id=teacher.id, name="Bad upload")
    db.add(module)
    db.commit()
    db.refresh(module)

    try:
        headers = _auth(client, "teacher@test.com", "password123")
        res = client.post(
            f"{MODULES_URL}/{module.id}/module-book",
            files={"file": ("notes.txt", b"plain text", "text/plain")},
            headers=headers,
        )
        assert res.status_code == 422
    finally:
        db.query(Module).filter(Module.id == module.id).delete()
        db.commit()


# ---------------------------------------------------------------------------
# G2-105 (slice a): re-parse + persist document text on upload/replace.
#
# The extracted text lives on FileRecord.extracted_text, which is intentionally
# NOT exposed by the API (RubricFileOut omits it), so these tests assert via a
# direct DB query rather than the response body.
# ---------------------------------------------------------------------------

import io  # noqa: E402

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
_XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _docx_bytes(paragraphs: list[str]) -> bytes:
    """Build a real, parseable .docx in memory."""
    from docx import Document

    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes(rows: list[list[str]]) -> bytes:
    """Build a real, parseable .xlsx in memory."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _record_for(db, module):
    from app.models.file_record import FileRecord

    db.refresh(module)
    return db.query(FileRecord).filter(FileRecord.id == module.module_book_id).first()


class TestModuleDocumentTextExtraction:
    def _module(self, db, teacher, name):
        from app.models.module import Module

        module = Module(id=uuid.uuid4(), teacher_id=teacher.id, name=name)
        db.add(module)
        db.commit()
        db.refresh(module)
        return module

    def _cleanup(self, db, module):
        from app.models.file_record import FileRecord
        from app.models.module import Module

        db.query(FileRecord).delete()
        db.query(Module).filter(Module.id == module.id).delete()
        db.commit()

    def test_module_book_docx_text_is_persisted(self, client, db, teacher):
        module = self._module(db, teacher, "Book with text")
        try:
            headers = _auth(client, "teacher@test.com", "password123")
            docx = _docx_bytes(["Criterion 1: database design", "Criterion 2: security"])
            up = client.post(
                f"{MODULES_URL}/{module.id}/module-book",
                files={"file": ("book.docx", docx, _DOCX_MIME)},
                headers=headers,
            )
            assert up.status_code == 200, up.text

            record = _record_for(db, module)
            assert record is not None
            assert record.extracted_text is not None
            assert "database design" in record.extracted_text
            assert "security" in record.extracted_text
        finally:
            self._cleanup(db, module)

    def test_rubric_xlsx_text_is_persisted(self, client, db, teacher):
        from app.models.file_record import FileRecord

        module = self._module(db, teacher, "Rubric with text")
        try:
            headers = _auth(client, "teacher@test.com", "password123")
            xlsx = _xlsx_bytes([["Criterion", "Weight"], ["database design", "40"]])
            up = client.post(
                f"{MODULES_URL}/{module.id}/rubric",
                files={"file": ("rubric.xlsx", xlsx, _XLSX_MIME)},
                headers=headers,
            )
            assert up.status_code == 200, up.text

            db.refresh(module)
            record = db.query(FileRecord).filter(
                FileRecord.id == module.rubric_file_id
            ).first()
            assert record is not None
            assert record.extracted_text is not None
            assert "database design" in record.extracted_text
        finally:
            self._cleanup(db, module)

    def test_replace_module_book_refreshes_text_and_drops_old(self, client, db, teacher):
        from app.models.file_record import FileRecord

        module = self._module(db, teacher, "Book replaced")
        try:
            headers = _auth(client, "teacher@test.com", "password123")

            up1 = client.post(
                f"{MODULES_URL}/{module.id}/module-book",
                files={"file": ("v1.docx", _docx_bytes(["alpha apple version one"]), _DOCX_MIME)},
                headers=headers,
            )
            assert up1.status_code == 200, up1.text

            up2 = client.post(
                f"{MODULES_URL}/{module.id}/module-book",
                files={"file": ("v2.docx", _docx_bytes(["bravo banana version two"]), _DOCX_MIME)},
                headers=headers,
            )
            assert up2.status_code == 200, up2.text

            # Only the new record survives — the old text is gone with it.
            assert db.query(FileRecord).count() == 1
            record = _record_for(db, module)
            assert "bravo banana" in record.extracted_text
            assert "alpha" not in record.extracted_text
        finally:
            self._cleanup(db, module)

    def test_unparseable_file_still_succeeds_with_null_text(self, client, db, teacher):
        """Non-fatal contract: a file that can't be parsed must NOT block the
        upload; it just stores no text. (Locks in the best-effort behaviour the
        existing junk-byte tests rely on incidentally.)"""
        module = self._module(db, teacher, "Book unparseable")
        try:
            headers = _auth(client, "teacher@test.com", "password123")
            up = client.post(
                f"{MODULES_URL}/{module.id}/module-book",
                files={"file": ("broken.pdf", b"%PDF-1.4 not a real pdf body", "application/pdf")},
                headers=headers,
            )
            assert up.status_code == 200, up.text

            record = _record_for(db, module)
            assert record is not None
            assert record.extracted_text is None
        finally:
            self._cleanup(db, module)


class TestModuleGithubRepos:
    def test_can_set_and_remove_student_repo(self, client, modules_teacher):
        headers = _auth(client)
        module = _create_module(client, headers, name="Repo Student")

        created = client.post(
            f"{MODULES_URL}/{module['id']}/students",
            json={"name": "Repo User", "student_number": "3001001"},
            headers=headers,
        )
        assert created.status_code == 201, created.text
        student_id = created.json()["id"]

        set_repo = client.patch(
            f"{MODULES_URL}/{module['id']}/students/{student_id}",
            json={"github_repo_url": "github.com/octocat/hello-world"},
            headers=headers,
        )
        assert set_repo.status_code == 200, set_repo.text
        assert set_repo.json()["github_repo_url"] == "https://github.com/octocat/hello-world"

        clear_repo = client.patch(
            f"{MODULES_URL}/{module['id']}/students/{student_id}",
            json={"github_repo_url": None},
            headers=headers,
        )
        assert clear_repo.status_code == 200, clear_repo.text
        assert clear_repo.json()["github_repo_url"] is None

    def test_group_repo_syncs_all_students(self, client, modules_teacher):
        headers = _auth(client)
        module = _create_module(client, headers, name="Repo Group")

        group_res = client.post(
            f"{MODULES_URL}/{module['id']}/groups",
            json={"name": "Group Repo"},
            headers=headers,
        )
        assert group_res.status_code == 201, group_res.text
        group_id = group_res.json()["id"]

        first = client.post(
            f"{MODULES_URL}/{module['id']}/students",
            json={"name": "One", "student_number": "3002001", "project_id": group_id},
            headers=headers,
        )
        assert first.status_code == 201, first.text

        second = client.post(
            f"{MODULES_URL}/{module['id']}/students",
            json={"name": "Two", "student_number": "3002002", "project_id": group_id},
            headers=headers,
        )
        assert second.status_code == 201, second.text

        set_group_repo = client.patch(
            f"{MODULES_URL}/{module['id']}/groups/{group_id}",
            json={"github_repo_url": "https://github.com/example/team-project"},
            headers=headers,
        )
        assert set_group_repo.status_code == 200, set_group_repo.text
        assert (
            set_group_repo.json()["github_repo_url"]
            == "https://github.com/example/team-project"
        )

        students = client.get(f"{MODULES_URL}/{module['id']}/students", headers=headers)
        assert students.status_code == 200, students.text
        repo_urls = {
            s["student_number"]: s["github_repo_url"] for s in students.json() if s["project_id"] == group_id
        }
        assert repo_urls["3002001"] == "https://github.com/example/team-project"
        assert repo_urls["3002002"] == "https://github.com/example/team-project"
# ---------------------------------------------------------------------------
# FR-03: serve module documents (rubric / module book) for in-app viewing
# ---------------------------------------------------------------------------


class TestServeModuleDocuments:
    def _module(self, db, teacher, name):
        from app.models.module import Module

        module = Module(id=uuid.uuid4(), teacher_id=teacher.id, name=name)
        db.add(module)
        db.commit()
        db.refresh(module)
        return module

    def _cleanup(self, db, module):
        from app.models.file_record import FileRecord
        from app.models.module import Module

        db.query(FileRecord).delete()
        db.query(Module).filter(Module.id == module.id).delete()
        db.commit()

    def test_view_rubric_pdf_returns_inline_bytes(self, client, db, teacher):
        module = self._module(db, teacher, "View rubric pdf")
        try:
            headers = _auth(client, "teacher@test.com", "password123")
            body = b"%PDF-1.4 viewable rubric"
            up = client.post(
                f"{MODULES_URL}/{module.id}/rubric",
                files={"file": ("rubric.pdf", body, "application/pdf")},
                headers=headers,
            )
            assert up.status_code == 200, up.text

            res = client.get(f"{MODULES_URL}/{module.id}/rubric/file", headers=headers)
            assert res.status_code == 200, res.text
            assert res.content == body
            assert res.headers["content-type"].startswith("application/pdf")
            assert "inline" in res.headers.get("content-disposition", "")
        finally:
            self._cleanup(db, module)

    def test_rubric_content_returns_extracted_text(self, client, db, teacher):
        module = self._module(db, teacher, "View rubric content")
        try:
            headers = _auth(client, "teacher@test.com", "password123")
            xlsx = _xlsx_bytes([["Criterion", "Weight"], ["database design", "40"]])
            up = client.post(
                f"{MODULES_URL}/{module.id}/rubric",
                files={"file": ("rubric.xlsx", xlsx, _XLSX_MIME)},
                headers=headers,
            )
            assert up.status_code == 200, up.text

            res = client.get(
                f"{MODULES_URL}/{module.id}/rubric/content", headers=headers
            )
            assert res.status_code == 200, res.text
            assert "database design" in res.json()["content"]
        finally:
            self._cleanup(db, module)

    def test_view_module_book_file_and_content(self, client, db, teacher):
        module = self._module(db, teacher, "View module book")
        try:
            headers = _auth(client, "teacher@test.com", "password123")
            docx = _docx_bytes(["Course intro", "Assessment outline"])
            up = client.post(
                f"{MODULES_URL}/{module.id}/module-book",
                files={"file": ("book.docx", docx, _DOCX_MIME)},
                headers=headers,
            )
            assert up.status_code == 200, up.text

            file_res = client.get(
                f"{MODULES_URL}/{module.id}/module-book/file", headers=headers
            )
            assert file_res.status_code == 200
            assert file_res.content == docx

            content_res = client.get(
                f"{MODULES_URL}/{module.id}/module-book/content", headers=headers
            )
            assert content_res.status_code == 200
            assert "Assessment outline" in content_res.json()["content"]
        finally:
            self._cleanup(db, module)

    def test_view_rubric_when_none_attached_returns_404(self, client, db, teacher):
        module = self._module(db, teacher, "No rubric")
        try:
            headers = _auth(client, "teacher@test.com", "password123")
            res = client.get(f"{MODULES_URL}/{module.id}/rubric/file", headers=headers)
            assert res.status_code == 404
        finally:
            self._cleanup(db, module)

    def test_view_other_teachers_rubric_returns_404(self, client, db, teacher):
        """A teacher must not be able to view a rubric on a module they don't own."""
        from app.models.module import Module
        from app.models.teacher import Teacher

        owner = Teacher(
            id=uuid.uuid4(),
            name="Doc Owner",
            email=f"doc_owner_{uuid.uuid4().hex[:8]}@test.com",
            password_hash=hash_password("password123"),
        )
        db.add(owner)
        db.commit()
        owner_module = Module(id=uuid.uuid4(), teacher_id=owner.id, name="Owner module")
        db.add(owner_module)
        db.commit()
        db.refresh(owner_module)

        # Owner uploads a rubric.
        owner_headers = _auth(client, owner.email, "password123")
        up = client.post(
            f"{MODULES_URL}/{owner_module.id}/rubric",
            files={"file": ("rubric.pdf", b"%PDF-1.4 private", "application/pdf")},
            headers=owner_headers,
        )
        assert up.status_code == 200, up.text

        try:
            # A different teacher must get 404 (module not visible to them).
            other_headers = _auth(client, "teacher@test.com", "password123")
            res = client.get(
                f"{MODULES_URL}/{owner_module.id}/rubric/file", headers=other_headers
            )
            assert res.status_code == 404
        finally:
            from app.models.file_record import FileRecord

            db.query(FileRecord).delete()
            db.query(Module).filter(Module.id == owner_module.id).delete()
            db.query(Teacher).filter(Teacher.id == owner.id).delete()
            db.commit()

    def test_viewing_document_writes_audit_event(self, client, db, teacher):
        from app.models.audit_event import AuditEvent

        module = self._module(db, teacher, "Audited view")
        try:
            headers = _auth(client, "teacher@test.com", "password123")
            up = client.post(
                f"{MODULES_URL}/{module.id}/rubric",
                files={"file": ("rubric.pdf", b"%PDF-1.4 audited", "application/pdf")},
                headers=headers,
            )
            assert up.status_code == 200, up.text

            before = (
                db.query(AuditEvent)
                .filter(AuditEvent.action == "document.viewed")
                .count()
            )
            res = client.get(f"{MODULES_URL}/{module.id}/rubric/file", headers=headers)
            assert res.status_code == 200

            events = (
                db.query(AuditEvent)
                .filter(AuditEvent.action == "document.viewed")
                .all()
            )
            assert len(events) == before + 1
            latest = events[-1]
            assert latest.details_json["kind"] == "rubric"
            assert latest.details_json["module_id"] == str(module.id)
        finally:
            db.query(AuditEvent).filter(
                AuditEvent.action == "document.viewed"
            ).delete()
            self._cleanup(db, module)


# ---------------------------------------------------------------------------
# Admin view-only: access control for modules owned by other teachers
# ---------------------------------------------------------------------------

_ADMIN_EMAIL = "admin_viewonly@test.com"
_ADMIN_PASSWORD = "adminpass123"
_OTHER_TEACHER_EMAIL = "other_teacher_viewonly@test.com"
_OTHER_TEACHER_PASSWORD = "otherpass123"


@pytest.fixture
def admin_teacher(db):
    """An admin teacher used by the admin view-only tests."""
    from app.models.teacher import Teacher

    existing = db.query(Teacher).filter(Teacher.email == _ADMIN_EMAIL).first()
    if existing:
        yield existing
        return

    t = Teacher(
        id=uuid.uuid4(),
        name="Admin Teacher",
        email=_ADMIN_EMAIL,
        password_hash=hash_password(_ADMIN_PASSWORD),
        is_admin=True,
    )
    db.add(t)
    db.commit()
    db.refresh(t)
    yield t


@pytest.fixture
def other_teacher(db):
    """A non-admin teacher who owns modules that the admin should only view."""
    from app.models.teacher import Teacher

    existing = db.query(Teacher).filter(Teacher.email == _OTHER_TEACHER_EMAIL).first()
    if existing:
        yield existing
        return

    t = Teacher(
        id=uuid.uuid4(),
        name="Other Teacher",
        email=_OTHER_TEACHER_EMAIL,
        password_hash=hash_password(_OTHER_TEACHER_PASSWORD),
    )
    db.add(t)
    db.commit()
    db.refresh(t)
    yield t


def _admin_auth(client) -> dict:
    res = client.post(LOGIN_URL, json={"email": _ADMIN_EMAIL, "password": _ADMIN_PASSWORD})
    assert res.status_code == 200, res.text
    return {"Authorization": f"Bearer {res.json()['access_token']}"}


def _create_module_for_teacher(db, teacher_id, name="Other Module"):
    from app.models.module import Module

    m = Module(id=uuid.uuid4(), teacher_id=teacher_id, name=name)
    db.add(m)
    db.commit()
    db.refresh(m)
    return m


class TestAdminViewOnly:
    """Admins can view but not mutate modules owned by other teachers."""

    def test_admin_can_list_all_modules(self, client, admin_teacher, other_teacher, db):
        m = _create_module_for_teacher(db, other_teacher.id, "Listed Module")
        try:
            res = client.get(MODULES_URL, headers=_admin_auth(client))
            assert res.status_code == 200
            ids = [item["id"] for item in res.json()]
            assert str(m.id) in ids
        finally:
            db.delete(m)
            db.commit()

    def test_admin_can_view_other_teachers_module(self, client, admin_teacher, other_teacher, db):
        m = _create_module_for_teacher(db, other_teacher.id, "Viewable Module")
        try:
            res = client.get(f"{MODULES_URL}/{m.id}", headers=_admin_auth(client))
            assert res.status_code == 200
            body = res.json()
            assert body["id"] == str(m.id)
            assert body["teacher_id"] == str(other_teacher.id)
        finally:
            db.delete(m)
            db.commit()

    def test_admin_cannot_rename_other_teachers_module(self, client, admin_teacher, other_teacher, db):
        m = _create_module_for_teacher(db, other_teacher.id, "Rename Target")
        try:
            res = client.patch(
                f"{MODULES_URL}/{m.id}",
                json={"name": "Hijacked"},
                headers=_admin_auth(client),
            )
            assert res.status_code == 403
        finally:
            db.delete(m)
            db.commit()

    def test_admin_cannot_delete_other_teachers_module(self, client, admin_teacher, other_teacher, db):
        m = _create_module_for_teacher(db, other_teacher.id, "Delete Target")
        try:
            res = client.delete(f"{MODULES_URL}/{m.id}", headers=_admin_auth(client))
            assert res.status_code == 403
        finally:
            db.delete(m)
            db.commit()

    def test_admin_cannot_create_group_in_other_teachers_module(
        self, client, admin_teacher, other_teacher, db
    ):
        m = _create_module_for_teacher(db, other_teacher.id, "Group Target")
        try:
            res = client.post(
                f"{MODULES_URL}/{m.id}/groups",
                json={"name": "Injected Group"},
                headers=_admin_auth(client),
            )
            assert res.status_code == 403
        finally:
            db.delete(m)
            db.commit()

    def test_admin_cannot_add_student_to_other_teachers_module(
        self, client, admin_teacher, other_teacher, db
    ):
        m = _create_module_for_teacher(db, other_teacher.id, "Student Target")
        try:
            res = client.post(
                f"{MODULES_URL}/{m.id}/students",
                json={"name": "Injected Student", "student_number": "9999999"},
                headers=_admin_auth(client),
            )
            assert res.status_code == 403
        finally:
            db.delete(m)
            db.commit()

    def test_admin_can_rename_own_module(self, client, admin_teacher, db):
        """Admin retains full permissions on their own modules."""
        m = _create_module_for_teacher(db, admin_teacher.id, "Admin Own Module")
        try:
            res = client.patch(
                f"{MODULES_URL}/{m.id}",
                json={"name": "Admin Renamed"},
                headers=_admin_auth(client),
            )
            assert res.status_code == 200
            assert res.json()["name"] == "Admin Renamed"
        finally:
            db.delete(m)
            db.commit()

    def test_admin_can_delete_own_module(self, client, admin_teacher, db):
        """Admin retains full permissions to delete their own modules."""
        m = _create_module_for_teacher(db, admin_teacher.id, "Admin Delete Me")
        res = client.delete(f"{MODULES_URL}/{m.id}", headers=_admin_auth(client))
        assert res.status_code == 204

    def test_module_out_includes_teacher_id(self, client, admin_teacher, other_teacher, db):
        """ModuleOut now exposes teacher_id so the frontend can detect view mode."""
        m = _create_module_for_teacher(db, other_teacher.id, "Teacher ID Module")
        try:
            res = client.get(f"{MODULES_URL}/{m.id}", headers=_admin_auth(client))
            assert res.status_code == 200
            assert "teacher_id" in res.json()
            assert res.json()["teacher_id"] == str(other_teacher.id)
        finally:
            db.delete(m)
            db.commit()

    def test_admin_cannot_import_students_to_other_teachers_module(
        self, client, admin_teacher, other_teacher, db
    ):
        import io

        m = _create_module_for_teacher(db, other_teacher.id, "Import Target")
        try:
            csv_content = b"Name,Student Number\nJohn Doe,1234567\n"
            res = client.post(
                f"{MODULES_URL}/{m.id}/students/import",
                files={"file": ("students.csv", io.BytesIO(csv_content), "text/csv")},
                headers=_admin_auth(client),
            )
            assert res.status_code == 403
        finally:
            db.delete(m)
            db.commit()
