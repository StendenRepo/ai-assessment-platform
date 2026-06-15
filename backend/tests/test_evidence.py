"""
Tests for the evidence upload user story.

Acceptance criteria:
  1. Teacher can upload a .md file as student evidence.
  2. File is stored and linked to the student.
  3. Content is read correctly.

The service is designed to be extensible: SUPPORTED_EXTENSIONS in
evidence_service.py controls which file types are accepted.
"""

import io
import uuid

import pytest
from PIL import Image

UPLOAD_URL = "/api/v1/students/{student_id}/evidence"
LIST_URL = "/api/v1/students/{student_id}/evidence"
CONTENT_URL = "/api/v1/evidence/{evidence_id}/content"
FILE_URL = "/api/v1/evidence/{evidence_id}/file"
LOGIN_URL = "/api/v1/auth/login"


# ── Helpers ───────────────────────────────────────────────────────────────────

def _auth_header(client, teacher):
    res = client.post(LOGIN_URL, json={"email": teacher.email, "password": "password123"})
    token = res.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def _make_md_file(content: str = "# Hello\n\nThis is evidence.", filename: str = "evidence.md"):
    return ("file", (filename, io.BytesIO(content.encode()), "text/markdown"))


def _make_png_file(filename: str = "evidence.png"):
    image = Image.new("RGB", (2, 2), color=(255, 255, 255))
    buffer = io.BytesIO()
    image.save(buffer, format="PNG")
    buffer.seek(0)
    return ("file", (filename, buffer, "image/png"))


def _evidence_root(tmp_path):
    return tmp_path / "evidence"


def _evidence_text_root(tmp_path):
    return tmp_path / "evidence_text"


@pytest.fixture
def student(db, teacher):
    """Create a minimal module + project + student linked to the teacher."""
    from app.models.module import Module
    from app.models.project import Project
    from app.models.student import Student

    module = Module(
        id=uuid.uuid4(),
        teacher_id=teacher.id,
        name="Test Module",
    )
    db.add(module)
    db.flush()

    project = Project(
        id=uuid.uuid4(),
        module_id=module.id,
        name="Test Project",
        group_name="Group A",
    )
    db.add(project)
    db.flush()

    s = Student(
        name="Alice",
        student_number="1001",
    )
    db.add(s)
    db.flush()
    s.projects.append(project)
    db.commit()
    db.refresh(s)
    yield s
    # Delete evidence first: tests upload Evidence rows that FK to this student.
    # Deleting the student while those rows exist raises IntegrityError, which
    # rolls back the teardown and leaves a committed teacher row behind —
    # cascading into UNIQUE(teachers.email) errors in later test files.
    from app.models.evidence import Evidence
    from app.models.student import student_projects

    db.query(Evidence).filter(Evidence.student_id == s.student_number).delete()
    db.execute(student_projects.delete().where(student_projects.c.student_id == s.student_number))
    db.delete(s)
    db.delete(project)
    db.delete(module)
    db.commit()


# ── AC 1: Teacher can upload a .md file as student evidence ───────────────────

class TestUploadMarkdownEvidence:
    def test_upload_md_returns_201(self, client, teacher, student):
        headers = _auth_header(client, teacher)
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        res = client.post(url, files=[_make_md_file()], headers=headers)
        assert res.status_code == 201

    def test_upload_md_response_contains_expected_fields(self, client, teacher, student):
        headers = _auth_header(client, teacher)
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        res = client.post(url, files=[_make_md_file()], headers=headers)
        body = res.json()
        assert body["file_name"] == "evidence.md"
        assert body["file_type"] == "markdown"
        assert body["source_type"] == "upload"
        assert body["embedding_status"] == "completed"

    def test_upload_unsupported_extension_returns_422(self, client, teacher, student):
        headers = _auth_header(client, teacher)
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        bad_file = ("file", ("archive.zip", io.BytesIO(b"PK\x03\x04"), "application/zip"))
        res = client.post(url, files=[bad_file], headers=headers)
        assert res.status_code == 422

    def test_upload_invalid_image_returns_422(self, client, teacher, student):
        headers = _auth_header(client, teacher)
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        bad_file = ("file", ("broken.png", io.BytesIO(b"not-an-image"), "image/png"))
        res = client.post(url, files=[bad_file], headers=headers)
        assert res.status_code == 422


class TestUploadImageEvidence:
    def test_upload_png_returns_201(self, client, teacher, student):
        headers = _auth_header(client, teacher)
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        res = client.post(url, files=[_make_png_file()], headers=headers)
        assert res.status_code == 201

    def test_upload_png_response_marks_image_type(self, client, teacher, student):
        headers = _auth_header(client, teacher)
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        res = client.post(url, files=[_make_png_file(filename="diagram.png")], headers=headers)
        body = res.json()
        assert body["file_name"] == "diagram.png"
        assert body["file_type"] == "image"
        assert body["source_type"] == "upload"
        assert body["embedding_status"] == "completed"

    def test_upload_without_auth_returns_401(self, client, student):
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        res = client.post(url, files=[_make_md_file()])
        assert res.status_code == 401

    def test_upload_unknown_student_returns_404(self, client, teacher):
        headers = _auth_header(client, teacher)
        url = UPLOAD_URL.format(student_id=str(uuid.uuid4()))
        res = client.post(url, files=[_make_md_file()], headers=headers)
        assert res.status_code == 404


# ── AC 2: File is stored and linked to the student ────────────────────────────

class TestEvidenceLinkedToStudent:
    def test_uploaded_evidence_appears_in_list(self, client, teacher, student):
        headers = _auth_header(client, teacher)
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        client.post(url, files=[_make_md_file(filename="linked.md")], headers=headers)

        list_url = LIST_URL.format(student_id=str(student.student_number))
        res = client.get(list_url, headers=headers)
        assert res.status_code == 200
        names = [e["file_name"] for e in res.json()]
        assert "linked.md" in names

    def test_evidence_student_id_matches(self, client, teacher, student):
        headers = _auth_header(client, teacher)
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        res = client.post(url, files=[_make_md_file()], headers=headers)
        assert res.json()["student_id"] == student.student_number

    def test_list_unknown_student_returns_404(self, client, teacher):
        headers = _auth_header(client, teacher)
        url = LIST_URL.format(student_id=str(uuid.uuid4()))
        res = client.get(url, headers=headers)
        assert res.status_code == 404

    def test_upload_preserves_raw_image_and_writes_text_sidecar(self, client, teacher, student, tmp_path, monkeypatch):
        monkeypatch.setattr("app.services.evidence_service.settings.UPLOAD_DIR", str(tmp_path))

        headers = _auth_header(client, teacher)
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        res = client.post(url, files=[_make_png_file(filename="proof.png")], headers=headers)
        assert res.status_code == 201

        saved_path = _evidence_root(tmp_path) / res.json()["file_path"]
        assert saved_path.exists()
        assert saved_path.read_bytes().startswith(b"\x89PNG")

        text_path = _evidence_text_root(tmp_path) / res.json()["file_path"]
        text_path = text_path.with_name(f"{text_path.name}.txt")
        assert text_path.exists()
        assert text_path.read_text(encoding="utf-8") == "[Image evidence uploaded: proof.png]"

    def test_upload_markdown_does_not_write_text_sidecar(self, client, teacher, student, tmp_path, monkeypatch):
        monkeypatch.setattr("app.services.evidence_service.settings.UPLOAD_DIR", str(tmp_path))

        headers = _auth_header(client, teacher)
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        res = client.post(url, files=[_make_md_file(filename="notes.md")], headers=headers)
        assert res.status_code == 201

        saved_path = _evidence_root(tmp_path) / res.json()["file_path"]
        assert saved_path.exists()
        assert saved_path.read_text(encoding="utf-8") == "# Hello\n\nThis is evidence."

        text_path = saved_path.with_name(f"{saved_path.name}.txt")
        assert not text_path.exists()


# ── AC 3: Content is read correctly ──────────────────────────────────────────

class TestReadEvidenceContent:
    def test_content_endpoint_returns_correct_text(self, client, teacher, student, tmp_path, monkeypatch):
        # Point UPLOAD_DIR to a temp directory so the file is actually written
        monkeypatch.setattr("app.services.evidence_service.settings.UPLOAD_DIR", str(tmp_path))

        headers = _auth_header(client, teacher)
        md_content = "# My Evidence\n\nThis is the content."
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        upload_res = client.post(
            url,
            files=[_make_md_file(content=md_content)],
            headers=headers,
        )
        assert upload_res.status_code == 201
        evidence_id = upload_res.json()["id"]

        content_url = CONTENT_URL.format(evidence_id=evidence_id)
        res = client.get(content_url, headers=headers)
        assert res.status_code == 200
        body = res.json()
        assert body["content"] == md_content
        assert body["file_name"] == "evidence.md"

    def test_content_unknown_evidence_returns_404(self, client, teacher):
        headers = _auth_header(client, teacher)
        url = CONTENT_URL.format(evidence_id=str(uuid.uuid4()))
        res = client.get(url, headers=headers)
        assert res.status_code == 404

    def test_image_content_endpoint_returns_fallback_text(self, client, teacher, student, tmp_path, monkeypatch):
        monkeypatch.setattr("app.services.evidence_service.settings.UPLOAD_DIR", str(tmp_path))
        # No vision model configured — should fall back to placeholder

        headers = _auth_header(client, teacher)
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        upload_res = client.post(
            url,
            files=[_make_png_file(filename="whiteboard.png")],
            headers=headers,
        )
        assert upload_res.status_code == 201
        evidence_id = upload_res.json()["id"]

        content_url = CONTENT_URL.format(evidence_id=evidence_id)
        res = client.get(content_url, headers=headers)
        assert res.status_code == 200
        body = res.json()
        assert body["file_name"] == "whiteboard.png"
        assert body["content"] == "[Image evidence uploaded: whiteboard.png]"

    def test_service_reprocess_rebuilds_text_from_stored_file(self, client, teacher, student, db, tmp_path, monkeypatch):
        monkeypatch.setattr("app.services.evidence_service.settings.UPLOAD_DIR", str(tmp_path))

        headers = _auth_header(client, teacher)
        upload_url = UPLOAD_URL.format(student_id=str(student.student_number))
        upload_res = client.post(
            upload_url,
            files=[_make_png_file(filename="board.png")],
            headers=headers,
        )
        assert upload_res.status_code == 201

        relative_path = upload_res.json()["file_path"]
        text_path = _evidence_text_root(tmp_path) / relative_path
        text_sidecar = text_path.with_name(f"{text_path.name}.txt")
        text_sidecar.parent.mkdir(parents=True, exist_ok=True)
        text_sidecar.write_text("stale text", encoding="utf-8")

        from app.services.evidence_service import EvidenceService

        _, content = EvidenceService.reprocess_content(upload_res.json()["id"], db)
        assert content == "[Image evidence uploaded: board.png]"
        assert text_sidecar.read_text(encoding="utf-8") == "[Image evidence uploaded: board.png]"

    def test_vision_model_text_is_saved_when_available(self, client, teacher, student, tmp_path, monkeypatch):
        monkeypatch.setattr("app.services.evidence_service.settings.UPLOAD_DIR", str(tmp_path))
        monkeypatch.setattr("app.services.evidence_service.settings.VISION_MODEL", "fake-vision")

        import httpx as _httpx

        class _FakeResponse:
            def raise_for_status(self):
                pass
            def json(self):
                return {"response": "A screenshot showing a student project dashboard."}

        class _FakeClient:
            def __enter__(self): return self
            def __exit__(self, *_): pass
            def post(self, *_args, **_kwargs): return _FakeResponse()

        monkeypatch.setattr(_httpx, "Client", lambda **_kw: _FakeClient())

        headers = _auth_header(client, teacher)
        upload_url = UPLOAD_URL.format(student_id=str(student.student_number))
        upload_res = client.post(
            upload_url,
            files=[_make_png_file(filename="dashboard.png")],
            headers=headers,
        )
        assert upload_res.status_code == 201

        content_url = CONTENT_URL.format(evidence_id=upload_res.json()["id"])
        content_res = client.get(content_url, headers=headers)
        assert content_res.status_code == 200
        assert content_res.json()["content"] == "A screenshot showing a student project dashboard."

    def test_content_read_retries_ai_when_sidecar_has_placeholder(self, client, teacher, student, tmp_path, monkeypatch):
        monkeypatch.setattr("app.services.evidence_service.settings.UPLOAD_DIR", str(tmp_path))
        monkeypatch.setattr("app.services.evidence_service.settings.VISION_MODEL", "")

        headers = _auth_header(client, teacher)
        upload_url = UPLOAD_URL.format(student_id=str(student.student_number))
        upload_res = client.post(
            upload_url,
            files=[_make_png_file(filename="retry.png")],
            headers=headers,
        )
        assert upload_res.status_code == 201

        # Simulate vision becoming available later; content read should refresh placeholder text.
        monkeypatch.setattr("app.services.evidence_service.settings.VISION_MODEL", "fake-vision")
        import httpx as _httpx

        class _FakeResponse:
            def raise_for_status(self):
                pass

            def json(self):
                return {"response": "Recovered AI text for retry image."}

        class _FakeClient:
            def __enter__(self):
                return self

            def __exit__(self, *_):
                pass

            def post(self, *_args, **_kwargs):
                return _FakeResponse()

        monkeypatch.setattr(_httpx, "Client", lambda **_kw: _FakeClient())

        content_url = CONTENT_URL.format(evidence_id=upload_res.json()["id"])
        content_res = client.get(content_url, headers=headers)
        assert content_res.status_code == 200
        assert content_res.json()["content"] == "Recovered AI text for retry image."

        sidecar = (_evidence_text_root(tmp_path) / upload_res.json()["file_path"]).with_name(
            f"{(_evidence_text_root(tmp_path) / upload_res.json()['file_path']).name}.txt"
        )
        assert sidecar.read_text(encoding="utf-8") == "Recovered AI text for retry image."


class TestReadEvidenceFile:
    def test_file_endpoint_returns_raw_image_bytes(self, client, teacher, student, tmp_path, monkeypatch):
        monkeypatch.setattr("app.services.evidence_service.settings.UPLOAD_DIR", str(tmp_path))

        headers = _auth_header(client, teacher)
        upload_url = UPLOAD_URL.format(student_id=str(student.student_number))
        upload_res = client.post(
            upload_url,
            files=[_make_png_file(filename="preview.png")],
            headers=headers,
        )
        assert upload_res.status_code == 201

        file_url = FILE_URL.format(evidence_id=upload_res.json()["id"])
        res = client.get(file_url, headers=headers)
        assert res.status_code == 200
        assert res.headers["content-type"] == "image/png"
        assert res.content.startswith(b"\x89PNG")

    def test_file_endpoint_requires_auth(self, client, teacher, student):
        headers = _auth_header(client, teacher)
        upload_url = UPLOAD_URL.format(student_id=str(student.student_number))
        upload_res = client.post(
            upload_url,
            files=[_make_png_file(filename="preview.png")],
            headers=headers,
        )
        assert upload_res.status_code == 201

        file_url = FILE_URL.format(evidence_id=upload_res.json()["id"])
        res = client.get(file_url)
        assert res.status_code == 401


# ── Extensibility: supported-types endpoint ───────────────────────────────────

class TestSupportedTypes:
    def test_supported_types_returns_list(self, client):
        res = client.get("/api/v1/evidence/supported-types")
        assert res.status_code == 200
        body = res.json()
        assert "supported_extensions" in body
        assert isinstance(body["supported_extensions"], list)

    def test_md_is_in_supported_types(self, client):
        res = client.get("/api/v1/evidence/supported-types")
        extensions = res.json()["supported_extensions"]
        assert ".md" in extensions
        assert ".pdf" in extensions
        assert ".docx" in extensions
        assert ".xlsx" in extensions
        assert ".csv" in extensions

    def test_png_is_in_supported_types(self, client):
        res = client.get("/api/v1/evidence/supported-types")
        assert ".png" in res.json()["supported_extensions"]

    def test_unsupported_extension_error_mentions_allowed_types(self, client, teacher, student):
        headers = _auth_header(client, teacher)
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        bad_file = ("file", ("archive.zip", io.BytesIO(b"PK\x03\x04"), "application/zip"))
        res = client.post(url, files=[bad_file], headers=headers)
        assert res.status_code == 422
        detail = res.json()["detail"]
        assert ".md" in detail
        assert ".pdf" in detail
        assert ".png" in detail


def _docx_bytes(paragraphs: list[str]) -> bytes:
    from docx import Document

    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes(rows: list[list[str]]) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestUploadOfficeEvidence:
    def test_upload_docx_extracts_text(self, client, teacher, student, tmp_path, monkeypatch):
        monkeypatch.setattr("app.services.evidence_service.settings.UPLOAD_DIR", str(tmp_path))
        headers = _auth_header(client, teacher)
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        docx = _docx_bytes(["Sprint report", "Shared authentication module notes"])
        res = client.post(
            url,
            files=[("file", ("report.docx", io.BytesIO(docx), "application/octet-stream"))],
            headers=headers,
        )
        assert res.status_code == 201
        assert res.json()["file_type"] == "docx"

        content_url = CONTENT_URL.format(evidence_id=res.json()["id"])
        body = client.get(content_url, headers=headers).json()
        assert "authentication module" in body["content"]

    def test_upload_xlsx_extracts_text(self, client, teacher, student, tmp_path, monkeypatch):
        monkeypatch.setattr("app.services.evidence_service.settings.UPLOAD_DIR", str(tmp_path))
        headers = _auth_header(client, teacher)
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        xlsx = _xlsx_bytes([["Task", "Status"], ["JWT middleware", "done"]])
        res = client.post(
            url,
            files=[("file", ("tasks.xlsx", io.BytesIO(xlsx), "application/octet-stream"))],
            headers=headers,
        )
        assert res.status_code == 201
        assert res.json()["file_type"] == "xlsx"

        content_url = CONTENT_URL.format(evidence_id=res.json()["id"])
        body = client.get(content_url, headers=headers).json()
        assert "JWT middleware" in body["content"]

    def test_upload_csv_extracts_text(self, client, teacher, student, tmp_path, monkeypatch):
        monkeypatch.setattr("app.services.evidence_service.settings.UPLOAD_DIR", str(tmp_path))
        headers = _auth_header(client, teacher)
        url = UPLOAD_URL.format(student_id=str(student.student_number))
        csv_data = b"Task,Status\nJWT middleware,done\n"
        res = client.post(
            url,
            files=[("file", ("tasks.csv", io.BytesIO(csv_data), "text/csv"))],
            headers=headers,
        )
        assert res.status_code == 201
        assert res.json()["file_type"] == "other"

        content_url = CONTENT_URL.format(evidence_id=res.json()["id"])
        body = client.get(content_url, headers=headers).json()
        assert "JWT middleware" in body["content"]


DELETE_URL = "/api/v1/evidence/{evidence_id}"


class TestDeleteEvidence:
    def test_delete_succeeds_when_overlap_signals_reference_evidence(
        self, client, teacher, student, db, tmp_path, monkeypatch
    ):
        from app.models.enums import OverlapType
        from app.models.overlap_signal import OverlapSignal

        monkeypatch.setattr("app.services.evidence_service.settings.UPLOAD_DIR", str(tmp_path))
        headers = _auth_header(client, teacher)
        upload_url = UPLOAD_URL.format(student_id=str(student.student_number))

        first = client.post(
            upload_url,
            files=[_make_md_file(content="Shared overlap text alpha", filename="a.md")],
            headers=headers,
        )
        second = client.post(
            upload_url,
            files=[_make_md_file(content="Shared overlap text alpha", filename="b.md")],
            headers=headers,
        )
        assert first.status_code == 201
        assert second.status_code == 201

        ev_a_id = uuid.UUID(first.json()["id"])
        ev_b_id = uuid.UUID(second.json()["id"])

        overlap = OverlapSignal(
            student_a_id=student.student_number,
            student_b_id=student.student_number,
            evidence_a_id=ev_a_id,
            evidence_b_id=ev_b_id,
            overlap_type=OverlapType.textual,
            confidence=0.7,
            snippet="overlap",
        )
        db.add(overlap)
        db.commit()
        overlap_id = overlap.id

        delete_url = DELETE_URL.format(evidence_id=str(ev_a_id))
        res = client.delete(delete_url, headers=headers)
        assert res.status_code == 204

        list_url = LIST_URL.format(student_id=str(student.student_number))
        remaining = client.get(list_url, headers=headers).json()
        assert all(item["id"] != str(ev_a_id) for item in remaining)
        assert (
            db.query(OverlapSignal).filter(OverlapSignal.id == overlap_id).first()
            is None
        )
