#!/usr/bin/env python3
"""Generate overlap-test-pack on the user's Desktop for manual QA."""

from __future__ import annotations

import csv
import io
import shutil
from pathlib import Path

from docx import Document
from openpyxl import Workbook

ROOT = Path.home() / "Desktop" / "overlap-test-pack"

SHARED_BLOCK = (
    "OVERLAP_MARKER_ALPHA: Our cohort implemented the authentication module using JWT tokens "
    "and bcrypt password hashing. We documented every REST endpoint in README.md and added "
    "integration tests for login, logout, and session refresh. The CLI planning tool persists "
    "tasks in SQLite behind a repository layer that isolates SQL from business rules."
)

UNIQUE_003 = (
    "UNIQUE_MARKER_GAMMA: Student three focused on Kanban board rendering, argparse subcommands, "
    "and CSV export for sprint retrospectives. No shared authentication boilerplate appears here."
)

UNIQUE_004 = (
    "UNIQUE_MARKER_DELTA: Student four implemented lint automation, pre-commit hooks, and weekly "
    "stand-up notes. This submission is intentionally independent from all other test students."
)


def write_md(path: Path, title: str, body: str, footer: str) -> None:
    path.write_text(f"# {title}\n\n{body}\n\n{footer}\n", encoding="utf-8")


def write_txt(path: Path, title: str, body: str, footer: str) -> None:
    path.write_text(f"{title}\n\n{body}\n\n{footer}\n", encoding="utf-8")


def write_docx(path: Path, title: str, body: str, footer: str) -> None:
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_paragraph(body)
    doc.add_paragraph(footer)
    doc.save(path)


def write_xlsx(path: Path, title: str, body: str, footer: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Submission"
    ws.append(["Section", "Content"])
    ws.append(["Title", title])
    ws.append(["Body", body])
    ws.append(["Footer", footer])
    buf = io.BytesIO()
    wb.save(buf)
    path.write_bytes(buf.getvalue())


def write_csv(path: Path, title: str, body: str, footer: str) -> None:
    with path.open("w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["section", "content"])
        writer.writerow(["title", title])
        writer.writerow(["body", body])
        writer.writerow(["footer", footer])


def write_pdf(path: Path, title: str, body: str, footer: str) -> None:
    from fpdf import FPDF

    def ascii_safe(value: str) -> str:
        return value.replace("\u2014", "-").replace("\u2013", "-")

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", size=11)
    pdf.multi_cell(0, 6, ascii_safe(f"{title}\n\n{body}\n\n{footer}"))
    pdf.output(str(path))


def student_set(base: Path, student_id: str, title: str, body: str, footer: str) -> None:
    folder = base / student_id
    folder.mkdir(parents=True, exist_ok=True)
    stem = "submission_pack"
    write_md(folder / f"{stem}.md", title, body, footer)
    write_txt(folder / f"{stem}.txt", title, body, footer)
    write_docx(folder / f"{stem}.docx", title, body, footer)
    write_xlsx(folder / f"{stem}.xlsx", title, body, footer)
    write_csv(folder / f"{stem}.csv", title, body, footer)
    write_pdf(folder / f"{stem}.pdf", title, body, footer)


def main() -> None:
    if ROOT.exists():
        shutil.rmtree(ROOT)
    ROOT.mkdir(parents=True)

    overlap_dir = ROOT / "01_should_detect_overlap"
    unique_dir = ROOT / "02_should_not_overlap"

    student_set(
        overlap_dir,
        "student-001",
        "Sprint 3 submission — Student 001",
        SHARED_BLOCK,
        "Student 001 led the repository refactor and pytest coverage for task CRUD.",
    )
    student_set(
        overlap_dir,
        "student-002",
        "Team deliverable — Student 002",
        SHARED_BLOCK,
        "Student 002 pair-programmed JWT middleware and updated endpoint examples.",
    )
    student_set(
        unique_dir,
        "student-003",
        "Individual reflection — Student 003",
        UNIQUE_003,
        "Owned persistence tests for duplicate titles and timezone-aware due dates.",
    )
    student_set(
        unique_dir,
        "student-004",
        "Weekly log — Student 004",
        UNIQUE_004,
        "Participated in stand-ups and fixed pylint warnings in the CLI module.",
    )

    readme = f"""# Overlap test pack

Manual overlap QA for the AI Assessment Platform.

## Folders

| Folder | Students | Overlap? |
|--------|----------|----------|
| `01_should_detect_overlap/` | student-001, student-002 | **Yes** — same `OVERLAP_MARKER_ALPHA` paragraph |
| `02_should_not_overlap/` | student-003, student-004 | **No** — unique markers GAMMA / DELTA |

Each student folder has **6 files**: `.md` `.txt` `.pdf` `.docx` `.xlsx` `.csv`

## Quick test

1. Log in as module owner (e.g. **bob.singh@university.edu** / **password123**).
2. Add students `9000101`–`9000104` (or use empty slots in a test group).
3. Upload files from each folder to the matching student page.
4. **Review overlaps** → **Scan module**.

## Expected

- **001 ↔ 002**: confirmed textual overlap (search for `OVERLAP_MARKER_ALPHA` in detail view).
- **003 ↔ 004**: no copy-paste overlap.
- **001/002 vs 003/004**: no strong textual overlap.

## Break-test ideas

- Upload all 6 formats for both overlap students → multiple signals (duplicates possible).
- Cross-format: 001 gets `.pdf`, 002 gets `.docx` → should still match on shared text.
- Filename-only noise: all use `submission_pack.*` — ignore 95% filename-only rows if body differs.

Shared block starts with:

`{SHARED_BLOCK[:100]}...`
"""
    (ROOT / "README.md").write_text(readme, encoding="utf-8")

    files = sorted(p for p in ROOT.rglob("*") if p.is_file())
    print(f"Created {ROOT} ({len(files)} files)")
    for p in files:
        print(f"  {p.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
